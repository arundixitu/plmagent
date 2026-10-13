import json
import hashlib
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter


# ==========================================================
# PROJECT PATHS
# ==========================================================

PROJECT_ROOT = Path(__file__).resolve().parent.parent

DATA_FOLDER = (
    PROJECT_ROOT
    / "data"
    / "raw"
)

EVALUATION_FOLDER = (
    PROJECT_ROOT
    / "evaluation"
)

PHASE5E_FOLDER = (
    EVALUATION_FOLDER
    / "phase5e"
)

PHASE5E_FOLDER.mkdir(
    parents=True,
    exist_ok=True
)


# ==========================================================
# SOURCE FILES
# ==========================================================

ENGINEERING_CHANGES_FILE = (
    DATA_FOLDER
    / "10_Engineering_Changes.xlsx"
)

COMPONENT_MASTER_FILE = (
    DATA_FOLDER
    / "04_Component_Master.xlsx"
)

REQUIREMENTS_FILE = (
    DATA_FOLDER
    / "06_Product_Requirements.xlsx"
)

VALIDATION_FILE = (
    DATA_FOLDER
    / "13_Test_Validation_Records.xlsx"
)

MANUFACTURING_FILE = (
    DATA_FOLDER
    / "12_Manufacturing_Records.xlsx"
)

SUPPLIER_FILE = (
    DATA_FOLDER
    / "09_Supplier_Records.xlsx"
)

SERVICE_FILE = (
    DATA_FOLDER
    / "14_Service_Records.xlsx"
)

DOCUMENT_FILE = (
    DATA_FOLDER
    / "16_Document_Control.xlsx"
)

RELATIONSHIP_MATRIX_FILE = (
    DATA_FOLDER
    / "18_Relationship_Matrix.xlsx"
)


# ==========================================================
# PROSPECTIVE UNSEEN CASES
# ==========================================================

UNSEEN_ECRS = [
    "ECR-002",
    "ECR-003",
    "ECR-010",
    "ECR-011",
    "ECR-014"
]


# ==========================================================
# OUTPUT FILES
# ==========================================================

LABELING_WORKBOOK = (
    PHASE5E_FOLDER
    / "Phase5E_Unseen_Expert_Labeling_Template.xlsx"
)

CASE_MANIFEST = (
    PHASE5E_FOLDER
    / "Phase5E_Unseen_Case_Manifest.json"
)


# ==========================================================
# HASHING
# ==========================================================

def calculate_sha256(
    file_path
):
    sha256 = hashlib.sha256()

    with open(
        file_path,
        "rb"
    ) as file:

        while True:

            chunk = file.read(
                8192
            )

            if not chunk:
                break

            sha256.update(
                chunk
            )

    return sha256.hexdigest()


# ==========================================================
# CLEAN DATAFRAME VALUES
# ==========================================================

def clean_value(
    value
):

    if pd.isna(
        value
    ):
        return ""

    return str(
        value
    )


# ==========================================================
# LOAD ENGINEERING CHANGES
# ==========================================================

def load_unseen_changes():

    if not ENGINEERING_CHANGES_FILE.exists():

        raise FileNotFoundError(
            (
                "Engineering change file not found:\n"
                f"{ENGINEERING_CHANGES_FILE}"
            )
        )


    df = pd.read_excel(
        ENGINEERING_CHANGES_FILE
    )


    if "ECR_ID" not in df.columns:

        raise ValueError(
            (
                "10_Engineering_Changes.xlsx does not "
                "contain ECR_ID."
            )
        )


    selected = df[
        df[
            "ECR_ID"
        ]
        .astype(str)
        .str.upper()
        .isin(
            UNSEEN_ECRS
        )
    ].copy()


    found = set(
        selected[
            "ECR_ID"
        ]
        .astype(str)
        .str.upper()
    )


    missing = (
        set(
            UNSEEN_ECRS
        )
        - found
    )


    if missing:

        raise ValueError(
            (
                "The following prospective ECRs "
                "were not found: "
                f"{sorted(missing)}"
            )
        )


    # Preserve requested order.
    selected[
        "_order"
    ] = selected[
        "ECR_ID"
    ].astype(str).str.upper().map(
        {
            ecr_id: index
            for index, ecr_id
            in enumerate(
                UNSEEN_ECRS
            )
        }
    )


    selected = selected.sort_values(
        "_order"
    ).drop(
        columns=[
            "_order"
        ]
    )


    return selected


# ==========================================================
# ADD SHEET FROM DATAFRAME
# ==========================================================

def add_dataframe_sheet(
    workbook,
    title,
    dataframe
):

    worksheet = workbook.create_sheet(
        title=title
    )


    # Header styling.
    header_fill = PatternFill(
        "solid",
        fgColor="D9EAF7"
    )

    thin_border = Border(
        left=Side(
            style="thin"
        ),
        right=Side(
            style="thin"
        ),
        top=Side(
            style="thin"
        ),
        bottom=Side(
            style="thin"
        )
    )


    for column_index, column_name in enumerate(
        dataframe.columns,
        start=1
    ):

        cell = worksheet.cell(
            row=1,
            column=column_index,
            value=str(
                column_name
            )
        )

        cell.font = Font(
            bold=True
        )

        cell.fill = header_fill

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

        cell.border = thin_border


    for row_index, row in enumerate(
        dataframe.itertuples(
            index=False
        ),
        start=2
    ):

        for column_index, value in enumerate(
            row,
            start=1
        ):

            cell = worksheet.cell(
                row=row_index,
                column=column_index,
                value=clean_value(
                    value
                )
            )

            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True
            )

            cell.border = thin_border


    # Column widths.
    for index, column_name in enumerate(
        dataframe.columns,
        start=1
    ):

        width = max(
            14,
            min(
                50,
                len(
                    str(
                        column_name
                    )
                )
                + 5
            )
        )

        worksheet.column_dimensions[
            get_column_letter(
                index
            )
        ].width = width


    worksheet.freeze_panes = "A2"

    return worksheet


# ==========================================================
# READ OPTIONAL REFERENCE FILE
# ==========================================================

def read_reference_file(
    file_path
):

    if not file_path.exists():
        return None

    try:

        return pd.read_excel(
            file_path
        )

    except Exception:

        return None


# ==========================================================
# BUILD EXPERT LABELING WORKBOOK
# ==========================================================

def build_labeling_workbook(
    changes_df
):

    workbook = Workbook()

    default_sheet = workbook.active

    workbook.remove(
        default_sheet
    )


    # ======================================================
    # SHEET 1 — INSTRUCTIONS
    # ======================================================

    instructions = workbook.create_sheet(
        "Instructions"
    )


    instructions[
        "A1"
    ] = "Phase 5E — Prospective Blind Expert Labeling"

    instructions[
        "A1"
    ].font = Font(
        bold=True,
        size=16
    )


    instructions[
        "A3"
    ] = (
        "Purpose"
    )

    instructions[
        "A3"
    ].font = Font(
        bold=True
    )


    instructions[
        "A4"
    ] = (
        "Create expert ground truth for five previously "
        "unlabelled engineering changes before D1 or D2.2 "
        "is evaluated on them."
    )


    instructions[
        "A6"
    ] = "Critical experimental rule"

    instructions[
        "A6"
    ].font = Font(
        bold=True
    )


    instructions[
        "A7"
    ] = (
        "Do NOT run D1, D2, D2.1 or D2.2 on these five "
        "ECRs until the expert labels have been completed "
        "and frozen."
    )


    instructions[
        "A9"
    ] = "Expert task"

    instructions[
        "A9"
    ].font = Font(
        bold=True
    )


    instructions[
        "A10"
    ] = (
        "For each ECR, identify all PLM entities that an "
        "engineering change-impact assessment should regard "
        "as genuinely impacted or requiring engineering "
        "disposition."
    )


    instructions[
        "A12"
    ] = "Include when supported"

    instructions[
        "A12"
    ].font = Font(
        bold=True
    )


    instructions[
        "A13"
    ] = (
        "Components, requirements, validation tests, "
        "manufacturing operations, suppliers, service "
        "records, or other PLM entities may be included "
        "when their engineering/lifecycle impact is "
        "material and evidence-supported."
    )


    instructions[
        "A15"
    ] = "Do not include merely because"

    instructions[
        "A15"
    ].font = Font(
        bold=True
    )


    instructions[
        "A16"
    ] = (
        "An entity exists, belongs to the same product, "
        "was retrieved, appears in a BOM, or is generally "
        "related to the subsystem."
    )


    instructions[
        "A18"
    ] = "Label format"

    instructions[
        "A18"
    ].font = Font(
        bold=True
    )


    instructions[
        "A19"
    ] = (
        "Enter impacted entity IDs separated by semicolons, "
        "for example: BAT-001; BMS-001; REQ-014; TEST-004"
    )


    instructions[
        "A21"
    ] = "Recommended expert process"

    instructions[
        "A21"
    ].font = Font(
        bold=True
    )


    instructions[
        "A22"
    ] = (
        "Preferably obtain independent labels from 2–3 "
        "PLM/mechanical/electrical engineering experts. "
        "Resolve disagreements by consensus before freezing "
        "the final benchmark."
    )


    instructions.column_dimensions[
        "A"
    ].width = 110


    for row in range(
        1,
        24
    ):

        instructions[
            f"A{row}"
        ].alignment = Alignment(
            vertical="top",
            wrap_text=True
        )


    # ======================================================
    # SHEET 2 — CASES TO LABEL
    # ======================================================

    label_columns = [

        "GT_ID",

        "ECR_ID",

        "Change_Title",

        "Primary_Component_ID",

        "Change_Description",

        "Recorded_Risk",

        "Status",

        "Expert_Expected_Impacted_Entities",

        "Expert_Risk_Level",

        "Expert_Rationale",

        "Reviewer_Name_or_Code",

        "Review_Status"
    ]


    label_rows = []


    for index, row in enumerate(
        changes_df.to_dict(
            orient="records"
        ),
        start=1
    ):

        label_rows.append(
            {

                "GT_ID":
                    f"NEWGT-{index:03d}",

                "ECR_ID":
                    clean_value(
                        row.get(
                            "ECR_ID"
                        )
                    ),

                "Change_Title":
                    clean_value(
                        row.get(
                            "Change_Title"
                        )
                    ),

                "Primary_Component_ID":
                    clean_value(
                        row.get(
                            "Primary_Component_ID"
                        )
                    ),

                "Change_Description":
                    clean_value(
                        row.get(
                            "Change_Description"
                        )
                    ),

                "Recorded_Risk":
                    clean_value(
                        row.get(
                            "Risk"
                        )
                    ),

                "Status":
                    clean_value(
                        row.get(
                            "Status"
                        )
                    ),

                "Expert_Expected_Impacted_Entities":
                    "",

                "Expert_Risk_Level":
                    "",

                "Expert_Rationale":
                    "",

                "Reviewer_Name_or_Code":
                    "",

                "Review_Status":
                    "Pending"
            }
        )


    labeling_df = pd.DataFrame(
        label_rows,
        columns=label_columns
    )


    label_sheet = add_dataframe_sheet(
        workbook,
        "Expert_Labeling",
        labeling_df
    )


    # Wider expert-entry columns.
    label_sheet.column_dimensions[
        "H"
    ].width = 70

    label_sheet.column_dimensions[
        "J"
    ].width = 70


    # ======================================================
    # REFERENCE SHEETS
    # ======================================================

    reference_sources = [

        (
            "Components",
            COMPONENT_MASTER_FILE
        ),

        (
            "Requirements",
            REQUIREMENTS_FILE
        ),

        (
            "Validation",
            VALIDATION_FILE
        ),

        (
            "Manufacturing",
            MANUFACTURING_FILE
        ),

        (
            "Suppliers",
            SUPPLIER_FILE
        ),

        (
            "Service",
            SERVICE_FILE
        ),

        (
            "Documents",
            DOCUMENT_FILE
        ),

        (
            "Relationships",
            RELATIONSHIP_MATRIX_FILE
        )
    ]


    for sheet_name, file_path in (
        reference_sources
    ):

        dataframe = read_reference_file(
            file_path
        )


        if dataframe is not None:

            add_dataframe_sheet(
                workbook,
                sheet_name,
                dataframe
            )


    workbook.save(
        LABELING_WORKBOOK
    )


# ==========================================================
# BUILD CASE MANIFEST
# ==========================================================

def build_manifest(
    changes_df
):

    source_hashes = {}


    source_files = [

        ENGINEERING_CHANGES_FILE,
        COMPONENT_MASTER_FILE,
        REQUIREMENTS_FILE,
        VALIDATION_FILE,
        MANUFACTURING_FILE,
        SUPPLIER_FILE,
        SERVICE_FILE,
        DOCUMENT_FILE,
        RELATIONSHIP_MATRIX_FILE
    ]


    for file_path in source_files:

        if file_path.exists():

            source_hashes[
                file_path.name
            ] = calculate_sha256(
                file_path
            )


    cases = []


    for row in changes_df.to_dict(
        orient="records"
    ):

        cases.append(
            {

                "ECR_ID":
                    clean_value(
                        row.get(
                            "ECR_ID"
                        )
                    ),

                "Change_Title":
                    clean_value(
                        row.get(
                            "Change_Title"
                        )
                    ),

                "Primary_Component_ID":
                    clean_value(
                        row.get(
                            "Primary_Component_ID"
                        )
                    ),

                "Change_Description":
                    clean_value(
                        row.get(
                            "Change_Description"
                        )
                    ),

                "Risk":
                    clean_value(
                        row.get(
                            "Risk"
                        )
                    ),

                "Status":
                    clean_value(
                        row.get(
                            "Status"
                        )
                    )
            }
        )


    manifest = {

        "benchmark_name":
            (
                "Aureon Phase 5E Prospective "
                "Unseen Evaluation Set"
            ),

        "created_at":
            datetime.now().isoformat(),

        "purpose":
            (
                "Prospective comparison of frozen D1 "
                "against frozen D2.2 after expert ground "
                "truth is created."
            ),

        "evaluation_ecrs":
            UNSEEN_ECRS,

        "number_of_cases":
            len(
                UNSEEN_ECRS
            ),

        "ground_truth_status":
            "NOT YET CREATED",

        "agent_execution_allowed":
            False,

        "cases":
            cases,

        "source_file_sha256":
            source_hashes
    }


    with open(
        CASE_MANIFEST,
        "w",
        encoding="utf-8"
    ) as file:

        json.dump(
            manifest,
            file,
            indent=2,
            ensure_ascii=False
        )


# ==========================================================
# MAIN
# ==========================================================

if __name__ == "__main__":

    print(
        "\n"
        + "=" * 90
    )

    print(
        "AUREON AI-PLM AGENT"
    )

    print(
        "Phase 5E.0 - Prospective Unseen Benchmark Setup"
    )

    print(
        "=" * 90
    )


    print(
        "\nProspective evaluation ECRs:"
    )


    for ecr_id in UNSEEN_ECRS:

        print(
            f"  - {ecr_id}"
        )


    changes_df = (
        load_unseen_changes()
    )


    print(
        f"\nCases located successfully: "
        f"{len(changes_df)}"
    )


    build_labeling_workbook(
        changes_df
    )


    build_manifest(
        changes_df
    )


    print(
        "\n"
        + "=" * 90
    )

    print(
        "PHASE 5E.0 FILES CREATED"
    )

    print(
        "=" * 90
    )


    print(
        "\nExpert labeling workbook:"
    )

    print(
        LABELING_WORKBOOK
    )


    print(
        "\nFrozen prospective case manifest:"
    )

    print(
        CASE_MANIFEST
    )


    print(
        "\nIMPORTANT:"
    )

    print(
        "Do NOT run D1 or D2.2 on these five cases yet."
    )


    print(
        "\nComplete and freeze the expert ground truth first."
    )


    print(
        "\n"
        + "=" * 90
    )

    print(
        "PHASE 5E.0 COMPLETE"
    )

    print(
        "=" * 90
    )