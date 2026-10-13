# ==========================================================
# AUREON AI-PLM AGENT
# Phase 5E.1A
#
# Multi-Reviewer Blind Expert Labeling Pack
#
# Creates:
#   Reviewer_A workbook
#   Reviewer_B workbook
#   Reviewer_C workbook
#   Consensus workbook
#   Reviewer-pack manifest
#
# IMPORTANT:
# No agent predictions are accessed.
# No ground truth is generated automatically.
# ==========================================================


import json
import hashlib
import shutil
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import (
    Font,
    PatternFill,
    Alignment,
    Border,
    Side
)
from openpyxl.worksheet.datavalidation import (
    DataValidation
)


# ==========================================================
# PROJECT PATHS
# ==========================================================

PROJECT_ROOT = (
    Path(__file__)
    .resolve()
    .parent
    .parent
)


PHASE5E_FOLDER = (
    PROJECT_ROOT
    / "evaluation"
    / "phase5e"
)


SOURCE_TEMPLATE = (
    PHASE5E_FOLDER
    / "Phase5E_Unseen_Expert_Labeling_Template.xlsx"
)


REVIEWER_FOLDER = (
    PHASE5E_FOLDER
    / "reviewers"
)


REVIEWER_FOLDER.mkdir(
    parents=True,
    exist_ok=True
)


# ==========================================================
# OUTPUT FILES
# ==========================================================

REVIEWER_FILES = {

    "Reviewer_A":
        REVIEWER_FOLDER
        / "Phase5E_Reviewer_A.xlsx",

    "Reviewer_B":
        REVIEWER_FOLDER
        / "Phase5E_Reviewer_B.xlsx",

    "Reviewer_C":
        REVIEWER_FOLDER
        / "Phase5E_Reviewer_C.xlsx"
}


CONSENSUS_FILE = (
    REVIEWER_FOLDER
    / "Phase5E_Consensus.xlsx"
)


MANIFEST_FILE = (
    REVIEWER_FOLDER
    / "Phase5E_Reviewer_Pack_Manifest.json"
)


PROTOCOL_FILE = (
    REVIEWER_FOLDER
    / "Phase5E_Reviewer_Protocol.txt"
)


# ==========================================================
# EXPECTED ECRs
# ==========================================================

EXPECTED_ECRS = [

    "ECR-002",
    "ECR-003",
    "ECR-010",
    "ECR-011",
    "ECR-014"
]


# ==========================================================
# LABELING COLUMNS
# ==========================================================

EXPECTED_ENTITY_COLUMN = (
    "Expert_Expected_Impacted_Entities"
)

EXPERT_RISK_COLUMN = (
    "Expert_Risk_Level"
)

RATIONALE_COLUMN = (
    "Expert_Rationale"
)

REVIEWER_COLUMN = (
    "Reviewer_Name_or_Code"
)

STATUS_COLUMN = (
    "Review_Status"
)


# ==========================================================
# SHA-256
# ==========================================================

def calculate_sha256(
    file_path
):

    sha256 = hashlib.sha256()


    with open(
        file_path,
        "rb"
    ) as file:

        while True:

            chunk = file.read(
                8192
            )


            if not chunk:
                break


            sha256.update(
                chunk
            )


    return sha256.hexdigest()


# ==========================================================
# NORMALIZE TEXT
# ==========================================================

def normalize_text(
    value
):

    if value is None:
        return ""


    return str(
        value
    ).strip()


# ==========================================================
# FIND HEADER COLUMNS
# ==========================================================

def get_header_map(
    worksheet
):

    header_map = {}


    for cell in worksheet[
        1
    ]:

        if cell.value is None:
            continue


        header_map[
            normalize_text(
                cell.value
            )
        ] = cell.column


    return header_map


# ==========================================================
# VALIDATE SOURCE TEMPLATE
# ==========================================================

def validate_source_template():

    if not SOURCE_TEMPLATE.exists():

        raise FileNotFoundError(
            (
                "Source labeling template not found:\n"
                f"{SOURCE_TEMPLATE}"
            )
        )


    workbook = load_workbook(
        SOURCE_TEMPLATE,
        data_only=False
    )


    if (
        "Expert_Labeling"
        not in workbook.sheetnames
    ):

        raise ValueError(
            (
                "Source workbook does not contain "
                "'Expert_Labeling' sheet."
            )
        )


    worksheet = workbook[
        "Expert_Labeling"
    ]


    header_map = get_header_map(
        worksheet
    )


    required_columns = [

        "ECR_ID",

        EXPECTED_ENTITY_COLUMN,

        EXPERT_RISK_COLUMN,

        RATIONALE_COLUMN,

        REVIEWER_COLUMN,

        STATUS_COLUMN
    ]


    missing_columns = [

        column
        for column in required_columns
        if column not in header_map
    ]


    if missing_columns:

        raise ValueError(
            (
                "Missing required columns in "
                "Expert_Labeling sheet:\n"
                f"{missing_columns}"
            )
        )


    ecr_column = header_map[
        "ECR_ID"
    ]


    found_ecrs = []


    for row in range(
        2,
        worksheet.max_row + 1
    ):

        value = normalize_text(
            worksheet.cell(
                row=row,
                column=ecr_column
            ).value
        ).upper()


        if value:

            found_ecrs.append(
                value
            )


    missing_ecrs = [

        ecr_id
        for ecr_id in EXPECTED_ECRS
        if ecr_id not in found_ecrs
    ]


    if missing_ecrs:

        raise ValueError(
            (
                "Source workbook is missing "
                "prospective benchmark cases:\n"
                f"{missing_ecrs}"
            )
        )


    workbook.close()


# ==========================================================
# FORMAT REVIEWER SHEET
# ==========================================================

def format_reviewer_sheet(
    worksheet,
    reviewer_code
):

    header_map = get_header_map(
        worksheet
    )


    expected_col = header_map[
        EXPECTED_ENTITY_COLUMN
    ]


    risk_col = header_map[
        EXPERT_RISK_COLUMN
    ]


    rationale_col = header_map[
        RATIONALE_COLUMN
    ]


    reviewer_col = header_map[
        REVIEWER_COLUMN
    ]


    status_col = header_map[
        STATUS_COLUMN
    ]


    # ------------------------------------------------------
    # COLORS
    # ------------------------------------------------------

    editable_fill = PatternFill(
        fill_type="solid",
        fgColor="FFF2CC"
    )


    locked_info_fill = PatternFill(
        fill_type="solid",
        fgColor="E7E6E6"
    )


    reviewer_fill = PatternFill(
        fill_type="solid",
        fgColor="DDEBF7"
    )


    completed_fill = PatternFill(
        fill_type="solid",
        fgColor="E2F0D9"
    )


    thin_border = Border(

        left=Side(
            style="thin",
            color="BFBFBF"
        ),

        right=Side(
            style="thin",
            color="BFBFBF"
        ),

        top=Side(
            style="thin",
            color="BFBFBF"
        ),

        bottom=Side(
            style="thin",
            color="BFBFBF"
        )
    )


    # ------------------------------------------------------
    # FILL ROWS
    # ------------------------------------------------------

    for row in range(
        2,
        worksheet.max_row + 1
    ):

        # Existing case-description information.
        for column in range(
            1,
            expected_col
        ):

            worksheet.cell(
                row=row,
                column=column
            ).fill = (
                locked_info_fill
            )


        # Blank independent expert fields.
        worksheet.cell(
            row=row,
            column=expected_col
        ).value = ""


        worksheet.cell(
            row=row,
            column=risk_col
        ).value = ""


        worksheet.cell(
            row=row,
            column=rationale_col
        ).value = ""


        worksheet.cell(
            row=row,
            column=reviewer_col
        ).value = reviewer_code


        worksheet.cell(
            row=row,
            column=status_col
        ).value = "Pending"


        # Editable expert cells.
        for column in [

            expected_col,
            risk_col,
            rationale_col,
            status_col
        ]:

            worksheet.cell(
                row=row,
                column=column
            ).fill = editable_fill


        worksheet.cell(
            row=row,
            column=reviewer_col
        ).fill = reviewer_fill


        # Formatting.
        for column in range(
            1,
            worksheet.max_column + 1
        ):

            cell = worksheet.cell(
                row=row,
                column=column
            )


            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True
            )


            cell.border = thin_border


    # ------------------------------------------------------
    # DATA VALIDATION — RISK
    # ------------------------------------------------------

    risk_validation = DataValidation(
        type="list",
        formula1=(
            '"Low,Medium,High,Critical"'
        ),
        allow_blank=True
    )


    risk_validation.error = (
        "Select Low, Medium, High or Critical."
    )


    risk_validation.errorTitle = (
        "Invalid risk level"
    )


    worksheet.add_data_validation(
        risk_validation
    )


    risk_validation.add(
        worksheet.cell(
            row=2,
            column=risk_col
        )
    )


    if worksheet.max_row > 2:

        risk_validation.add(
            (
                f"{worksheet.cell(row=2, column=risk_col).coordinate}:"
                f"{worksheet.cell(row=worksheet.max_row, column=risk_col).coordinate}"
            )
        )


    # ------------------------------------------------------
    # DATA VALIDATION — STATUS
    # ------------------------------------------------------

    status_validation = DataValidation(
        type="list",
        formula1=(
            '"Pending,In Progress,Completed"'
        ),
        allow_blank=False
    )


    status_validation.error = (
        "Select Pending, In Progress or Completed."
    )


    status_validation.errorTitle = (
        "Invalid review status"
    )


    worksheet.add_data_validation(
        status_validation
    )


    status_validation.add(
        (
            f"{worksheet.cell(row=2, column=status_col).coordinate}:"
            f"{worksheet.cell(row=worksheet.max_row, column=status_col).coordinate}"
        )
    )


    # ------------------------------------------------------
    # CONDITIONAL FORMATTING
    # ------------------------------------------------------

    status_letter = (
        worksheet.cell(
            row=1,
            column=status_col
        ).column_letter
    )


    worksheet.conditional_formatting.add(

        (
            f"{status_letter}2:"
            f"{status_letter}{worksheet.max_row}"
        ),

        (
            __import__(
                "openpyxl"
            )
            .formatting.rule.CellIsRule(
                operator="equal",
                formula=['"Completed"'],
                fill=completed_fill
            )
        )
    )


    # ------------------------------------------------------
    # COLUMN WIDTHS
    # ------------------------------------------------------

    worksheet.column_dimensions[
        worksheet.cell(
            row=1,
            column=expected_col
        ).column_letter
    ].width = 70


    worksheet.column_dimensions[
        worksheet.cell(
            row=1,
            column=rationale_col
        ).column_letter
    ].width = 70


    worksheet.column_dimensions[
        worksheet.cell(
            row=1,
            column=reviewer_col
        ).column_letter
    ].width = 24


    worksheet.column_dimensions[
        worksheet.cell(
            row=1,
            column=status_col
        ).column_letter
    ].width = 18


    worksheet.freeze_panes = "A2"


# ==========================================================
# UPDATE INSTRUCTION SHEET
# ==========================================================

def update_instruction_sheet(
    workbook,
    reviewer_code
):

    if (
        "Instructions"
        not in workbook.sheetnames
    ):

        return


    worksheet = workbook[
        "Instructions"
    ]


    worksheet[
        "A24"
    ] = "Assigned reviewer"


    worksheet[
        "A24"
    ].font = Font(
        bold=True
    )


    worksheet[
        "A25"
    ] = reviewer_code


    worksheet[
        "A25"
    ].font = Font(
        bold=True,
        size=13
    )


    worksheet[
        "A27"
    ] = (
        "BLIND REVIEW RULE: Do not inspect another "
        "reviewer's workbook, agent output, prediction "
        "folder, benchmark result, or previous ground "
        "truth while completing this workbook."
    )


    worksheet[
        "A27"
    ].font = Font(
        bold=True
    )


    worksheet[
        "A29"
    ] = (
        "Set Review_Status to Completed only after "
        "all impacted entities, risk level and rationale "
        "have been independently reviewed."
    )


    worksheet[
        "A31"
    ] = (
        "Do not try to match a desired number of impacts. "
        "Include only entities supported by engineering "
        "judgment and the PLM evidence supplied in this "
        "workbook."
    )


    worksheet.column_dimensions[
        "A"
    ].width = 110


    for row in [

        24,
        25,
        27,
        29,
        31
    ]:

        worksheet[
            f"A{row}"
        ].alignment = Alignment(
            vertical="top",
            wrap_text=True
        )


# ==========================================================
# CREATE REVIEWER WORKBOOK
# ==========================================================

def create_reviewer_workbook(
    reviewer_code,
    output_file
):

    # Start from an identical source workbook.
    shutil.copy2(
        SOURCE_TEMPLATE,
        output_file
    )


    workbook = load_workbook(
        output_file
    )


    worksheet = workbook[
        "Expert_Labeling"
    ]


    format_reviewer_sheet(
        worksheet,
        reviewer_code
    )


    update_instruction_sheet(
        workbook,
        reviewer_code
    )


    workbook.save(
        output_file
    )


    workbook.close()


# ==========================================================
# CREATE CONSENSUS WORKBOOK
# ==========================================================

def create_consensus_workbook():

    shutil.copy2(
        SOURCE_TEMPLATE,
        CONSENSUS_FILE
    )


    workbook = load_workbook(
        CONSENSUS_FILE
    )


    worksheet = workbook[
        "Expert_Labeling"
    ]


    header_map = get_header_map(
        worksheet
    )


    expected_col = header_map[
        EXPECTED_ENTITY_COLUMN
    ]


    risk_col = header_map[
        EXPERT_RISK_COLUMN
    ]


    rationale_col = header_map[
        RATIONALE_COLUMN
    ]


    reviewer_col = header_map[
        REVIEWER_COLUMN
    ]


    status_col = header_map[
        STATUS_COLUMN
    ]


    consensus_fill = PatternFill(
        fill_type="solid",
        fgColor="E2F0D9"
    )


    for row in range(
        2,
        worksheet.max_row + 1
    ):

        worksheet.cell(
            row=row,
            column=expected_col
        ).value = ""


        worksheet.cell(
            row=row,
            column=risk_col
        ).value = ""


        worksheet.cell(
            row=row,
            column=rationale_col
        ).value = ""


        worksheet.cell(
            row=row,
            column=reviewer_col
        ).value = (
            "Consensus Panel"
        )


        worksheet.cell(
            row=row,
            column=status_col
        ).value = (
            "Not Adjudicated"
        )


        for column in [

            expected_col,
            risk_col,
            rationale_col,
            reviewer_col,
            status_col
        ]:

            worksheet.cell(
                row=row,
                column=column
            ).fill = consensus_fill


            worksheet.cell(
                row=row,
                column=column
            ).alignment = Alignment(
                vertical="top",
                wrap_text=True
            )


    if (
        "Instructions"
        in workbook.sheetnames
    ):

        instructions = workbook[
            "Instructions"
        ]


        instructions[
            "A24"
        ] = (
            "CONSENSUS WORKBOOK"
        )


        instructions[
            "A24"
        ].font = Font(
            bold=True,
            size=14
        )


        instructions[
            "A25"
        ] = (
            "Do not complete this workbook until "
            "Reviewer A, Reviewer B and Reviewer C "
            "have finished their independent reviews."
        )


        instructions[
            "A27"
        ] = (
            "Use this workbook only to record the final "
            "adjudicated expert ground truth after "
            "reviewer disagreement has been examined."
        )


        instructions[
            "A29"
        ] = (
            "The final consensus labels must be frozen "
            "before D1 or D2.2 is run on any Phase 5E ECR."
        )


        for row in [

            24,
            25,
            27,
            29
        ]:

            instructions[
                f"A{row}"
            ].alignment = Alignment(
                vertical="top",
                wrap_text=True
            )


    workbook.save(
        CONSENSUS_FILE
    )


    workbook.close()


# ==========================================================
# CREATE REVIEW PROTOCOL
# ==========================================================

def create_protocol_file():

    protocol = f"""
AUREON AI-PLM AGENT
PHASE 5E.1A — BLIND MULTI-REVIEWER PROTOCOL

Created:
{datetime.now().isoformat()}

Prospective unseen cases:
{", ".join(EXPECTED_ECRS)}

============================================================
OBJECTIVE
============================================================

Create independent expert ground-truth labels for the
prospective Phase 5E benchmark before any D1 or D2.2
prediction is generated.

============================================================
REVIEWERS
============================================================

Reviewer A:
Phase5E_Reviewer_A.xlsx

Reviewer B:
Phase5E_Reviewer_B.xlsx

Reviewer C:
Phase5E_Reviewer_C.xlsx

Each reviewer must work independently.

============================================================
BLINDING RULES
============================================================

While performing the independent review, a reviewer must
not inspect:

- another reviewer's workbook;
- D1 predictions;
- D2 predictions;
- D2.1 predictions;
- D2.2 predictions;
- evaluation/predictions folders;
- previous benchmark prediction outputs;
- ground-truth files for other experiments when those
  reveal answers to these cases.

The agent must NOT be run on:
{", ".join(EXPECTED_ECRS)}

until consensus ground truth has been frozen.

============================================================
LABELING RULE
============================================================

For each ECR, identify PLM entities that should genuinely
be regarded as impacted or requiring engineering
disposition because of the proposed change.

Record entity IDs as semicolon-separated values.

Example formatting only:

BAT-001; BMS-001; REQ-014; TEST-004

Do NOT use the example as guidance for any Phase 5E case.

============================================================
DO NOT INCLUDE AN ENTITY MERELY BECAUSE
============================================================

- it exists;
- it belongs to the product;
- it is in the BOM;
- it is nearby in the data;
- a supplier exists;
- a test exists;
- a service record exists.

There must be an engineering or lifecycle impact path.

============================================================
RISK
============================================================

Select:

Low
Medium
High
Critical

Risk should represent expert engineering judgement for the
change as a whole.

============================================================
RATIONALE
============================================================

Briefly explain the engineering path behind the selected
entities.

The rationale is especially important when including
second-order or lifecycle impacts.

============================================================
COMPLETION
============================================================

Set each case's Review_Status to:

Completed

only after the reviewer has completed that case.

============================================================
CONSENSUS
============================================================

Do not edit Phase5E_Consensus.xlsx until all independent
reviewer workbooks have been completed.

Agreement analysis will be performed before the panel
resolves disagreements.

The final consensus benchmark will then be frozen and
cryptographically hashed before agent evaluation.
"""


    with open(
        PROTOCOL_FILE,
        "w",
        encoding="utf-8"
    ) as file:

        file.write(
            protocol.strip()
        )


# ==========================================================
# CREATE MANIFEST
# ==========================================================

def create_manifest():

    reviewer_hashes = {}


    for reviewer_code, file_path in (
        REVIEWER_FILES.items()
    ):

        reviewer_hashes[
            reviewer_code
        ] = {

            "file":
                file_path.name,

            "sha256_initial":
                calculate_sha256(
                    file_path
                )
        }


    manifest = {

        "phase":
            "5E.1A",

        "name":
            (
                "Prospective Multi-Reviewer "
                "Blind Labeling Pack"
            ),

        "created_at":
            datetime.now().isoformat(),

        "source_template":
            str(
                SOURCE_TEMPLATE
            ),

        "source_template_sha256":
            calculate_sha256(
                SOURCE_TEMPLATE
            ),

        "prospective_ecrs":
            EXPECTED_ECRS,

        "reviewers": [
            "Reviewer_A",
            "Reviewer_B",
            "Reviewer_C"
        ],

        "reviewer_files":
            reviewer_hashes,

        "consensus_file":
            CONSENSUS_FILE.name,

        "consensus_sha256_initial":
            calculate_sha256(
                CONSENSUS_FILE
            ),

        "protocol_file":
            PROTOCOL_FILE.name,

        "independent_review_status":
            "NOT STARTED",

        "consensus_status":
            "NOT STARTED",

        "agent_evaluation_allowed":
            False,

        "next_phase":
            (
                "5E.1B - Inter-rater "
                "Agreement Analysis"
            )
    }


    with open(
        MANIFEST_FILE,
        "w",
        encoding="utf-8"
    ) as file:

        json.dump(
            manifest,
            file,
            indent=2,
            ensure_ascii=False
        )


# ==========================================================
# MAIN
# ==========================================================

if __name__ == "__main__":

    print(
        "\n"
        + "=" * 90
    )


    print(
        "AUREON AI-PLM AGENT"
    )


    print(
        "Phase 5E.1A - Multi-Reviewer Blind Labeling Pack"
    )


    print(
        "=" * 90
    )


    # ------------------------------------------------------
    # VALIDATE SOURCE
    # ------------------------------------------------------

    print(
        "\nValidating prospective labeling template..."
    )


    validate_source_template()


    print(
        "Template validation PASSED."
    )


    # ------------------------------------------------------
    # CREATE REVIEWER FILES
    # ------------------------------------------------------

    print(
        "\nCreating independent reviewer workbooks..."
    )


    for reviewer_code, output_file in (
        REVIEWER_FILES.items()
    ):

        create_reviewer_workbook(
            reviewer_code,
            output_file
        )


        print(
            f"  Created {reviewer_code}:"
        )


        print(
            f"    {output_file}"
        )


    # ------------------------------------------------------
    # CONSENSUS
    # ------------------------------------------------------

    print(
        "\nCreating consensus workbook..."
    )


    create_consensus_workbook()


    print(
        f"  {CONSENSUS_FILE}"
    )


    # ------------------------------------------------------
    # PROTOCOL
    # ------------------------------------------------------

    print(
        "\nCreating reviewer protocol..."
    )


    create_protocol_file()


    print(
        f"  {PROTOCOL_FILE}"
    )


    # ------------------------------------------------------
    # MANIFEST
    # ------------------------------------------------------

    print(
        "\nCreating reviewer-pack manifest..."
    )


    create_manifest()


    print(
        f"  {MANIFEST_FILE}"
    )


    # ------------------------------------------------------
    # FINAL
    # ------------------------------------------------------

    print(
        "\n"
        + "=" * 90
    )


    print(
        "PHASE 5E.1A COMPLETE"
    )


    print(
        "=" * 90
    )


    print(
        "\nIndependent reviewer workbooks:"
    )


    for file_path in (
        REVIEWER_FILES.values()
    ):

        print(
            f"  {file_path}"
        )


    print(
        "\nConsensus workbook:"
    )


    print(
        f"  {CONSENSUS_FILE}"
    )


    print(
        "\nExperimental rule:"
    )


    print(
        "  Do NOT run D1 or D2.2 on ECR-002, ECR-003, "
        "ECR-010, ECR-011 or ECR-014 yet."
    )


    print(
        "\nNext:"
    )


    print(
        "  Have Reviewer A, B and C complete their "
        "workbooks independently."
    )


    print(
        "  Do NOT complete the consensus workbook yet."
    )


    print(
        "  After all three reviews are finished, "
        "Phase 5E.1B will calculate inter-rater agreement."
    )


    print(
        "\n"
        + "=" * 90
    )