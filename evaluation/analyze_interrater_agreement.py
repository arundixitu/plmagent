# ==========================================================
# AUREON AI-PLM AGENT
# Phase 5E.1B
#
# Inter-Rater Agreement Analysis
#
# Reads:
#   Phase5E_Reviewer_A.xlsx
#   Phase5E_Reviewer_B.xlsx
#   Phase5E_Reviewer_C.xlsx
#
# Produces:
#   Inter-rater agreement JSON
#   Summary CSV
#   Entity-level disagreement CSV
#   Adjudication workbook
#
# IMPORTANT:
# This script does NOT access agent predictions.
# This script does NOT create final consensus ground truth.
# ==========================================================


import csv
import json
import statistics
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook, Workbook
from openpyxl.styles import (
    Font,
    PatternFill,
    Alignment,
    Border,
    Side
)


# ==========================================================
# PROJECT PATHS
# ==========================================================

PROJECT_ROOT = (
    Path(__file__)
    .resolve()
    .parent
    .parent
)


PHASE5E_FOLDER = (
    PROJECT_ROOT
    / "evaluation"
    / "phase5e"
)


REVIEWER_FOLDER = (
    PHASE5E_FOLDER
    / "reviewers"
)


RESULTS_FOLDER = (
    PHASE5E_FOLDER
    / "agreement_results"
)


RESULTS_FOLDER.mkdir(
    parents=True,
    exist_ok=True
)


# ==========================================================
# INPUT FILES
# ==========================================================

REVIEWER_FILES = {

    "Reviewer_A":
        REVIEWER_FOLDER
        / "Phase5E_Reviewer_A.xlsx",

    "Reviewer_B":
        REVIEWER_FOLDER
        / "Phase5E_Reviewer_B.xlsx",

    "Reviewer_C":
        REVIEWER_FOLDER
        / "Phase5E_Reviewer_C.xlsx"
}


CONSENSUS_TEMPLATE = (
    REVIEWER_FOLDER
    / "Phase5E_Consensus.xlsx"
)


# ==========================================================
# EXPECTED CASES
# ==========================================================

EXPECTED_ECRS = [

    "ECR-002",
    "ECR-003",
    "ECR-010",
    "ECR-011",
    "ECR-014"
]


# ==========================================================
# COLUMN NAMES
# ==========================================================

ENTITY_COLUMN = (
    "Expert_Expected_Impacted_Entities"
)

RISK_COLUMN = (
    "Expert_Risk_Level"
)

RATIONALE_COLUMN = (
    "Expert_Rationale"
)

STATUS_COLUMN = (
    "Review_Status"
)


# ==========================================================
# OUTPUT NAMES
# ==========================================================

TIMESTAMP = datetime.now().strftime(
    "%Y%m%d_%H%M%S"
)


SUMMARY_CSV = (
    RESULTS_FOLDER
    / f"Phase5E_Reviewer_Agreement_Summary_{TIMESTAMP}.csv"
)


ENTITY_CSV = (
    RESULTS_FOLDER
    / f"Phase5E_Entity_Disagreement_{TIMESTAMP}.csv"
)


ANALYSIS_JSON = (
    RESULTS_FOLDER
    / f"Phase5E_InterRater_Agreement_{TIMESTAMP}.json"
)


ADJUDICATION_WORKBOOK = (
    RESULTS_FOLDER
    / f"Phase5E_Adjudication_Workbook_{TIMESTAMP}.xlsx"
)


# ==========================================================
# NORMALIZATION
# ==========================================================

def normalize_text(
    value
):

    if value is None:
        return ""

    return str(
        value
    ).strip()


def normalize_entity_id(
    value
):

    return (
        normalize_text(
            value
        )
        .upper()
    )


# ==========================================================
# PARSE ENTITY LIST
# ==========================================================

def parse_entity_list(
    value
):
    """
    Convert semicolon/comma/newline-separated text
    into a normalized set of entity IDs.
    """

    if value is None:
        return set()


    text = str(
        value
    )


    # Normalize separators.
    for separator in [
        "\n",
        ",",
        "|"
    ]:

        text = text.replace(
            separator,
            ";"
        )


    entities = set()


    for item in text.split(
        ";"
    ):

        entity = normalize_entity_id(
            item
        )


        if entity:

            entities.add(
                entity
            )


    return entities


# ==========================================================
# HEADER MAP
# ==========================================================

def get_header_map(
    worksheet
):

    header_map = {}


    for cell in worksheet[
        1
    ]:

        if cell.value is None:
            continue


        header_map[
            normalize_text(
                cell.value
            )
        ] = cell.column


    return header_map


# ==========================================================
# LOAD ONE REVIEWER
# ==========================================================

def load_reviewer(
    reviewer_name,
    file_path
):

    if not file_path.exists():

        raise FileNotFoundError(
            (
                f"{reviewer_name} workbook not found:\n"
                f"{file_path}"
            )
        )


    workbook = load_workbook(
        file_path,
        data_only=False
    )


    if (
        "Expert_Labeling"
        not in workbook.sheetnames
    ):

        raise ValueError(
            (
                f"{reviewer_name} workbook does not contain "
                "'Expert_Labeling'."
            )
        )


    worksheet = workbook[
        "Expert_Labeling"
    ]


    header_map = get_header_map(
        worksheet
    )


    required_columns = [

        "ECR_ID",
        ENTITY_COLUMN,
        RISK_COLUMN,
        RATIONALE_COLUMN,
        STATUS_COLUMN
    ]


    missing = [

        column
        for column in required_columns
        if column not in header_map
    ]


    if missing:

        raise ValueError(
            (
                f"{reviewer_name} missing columns: "
                f"{missing}"
            )
        )


    ecr_col = header_map[
        "ECR_ID"
    ]


    entity_col = header_map[
        ENTITY_COLUMN
    ]


    risk_col = header_map[
        RISK_COLUMN
    ]


    rationale_col = header_map[
        RATIONALE_COLUMN
    ]


    status_col = header_map[
        STATUS_COLUMN
    ]


    results = {}


    for row in range(
        2,
        worksheet.max_row + 1
    ):

        ecr_id = normalize_entity_id(
            worksheet.cell(
                row=row,
                column=ecr_col
            ).value
        )


        if not ecr_id:
            continue


        if ecr_id not in EXPECTED_ECRS:
            continue


        results[
            ecr_id
        ] = {

            "entities":
                parse_entity_list(
                    worksheet.cell(
                        row=row,
                        column=entity_col
                    ).value
                ),

            "risk":
                normalize_text(
                    worksheet.cell(
                        row=row,
                        column=risk_col
                    ).value
                ),

            "rationale":
                normalize_text(
                    worksheet.cell(
                        row=row,
                        column=rationale_col
                    ).value
                ),

            "status":
                normalize_text(
                    worksheet.cell(
                        row=row,
                        column=status_col
                    ).value
                )
        }


    workbook.close()


    missing_cases = [

        ecr_id
        for ecr_id in EXPECTED_ECRS
        if ecr_id not in results
    ]


    if missing_cases:

        raise ValueError(
            (
                f"{reviewer_name} missing ECR cases: "
                f"{missing_cases}"
            )
        )


    return results


# ==========================================================
# COMPLETION VALIDATION
# ==========================================================

def validate_reviews_complete(
    all_reviews
):

    errors = []


    for reviewer_name, reviews in (
        all_reviews.items()
    ):

        for ecr_id in EXPECTED_ECRS:

            record = reviews[
                ecr_id
            ]


            status = (
                record[
                    "status"
                ]
                .strip()
                .lower()
            )


            if status != "completed":

                errors.append(
                    (
                        f"{reviewer_name} / {ecr_id}: "
                        f"Review_Status is "
                        f"'{record['status']}' "
                        "instead of 'Completed'."
                    )
                )


            if not record[
                "entities"
            ]:

                errors.append(
                    (
                        f"{reviewer_name} / {ecr_id}: "
                        "no impacted entities entered."
                    )
                )


            if not record[
                "risk"
            ]:

                errors.append(
                    (
                        f"{reviewer_name} / {ecr_id}: "
                        "risk level is blank."
                    )
                )


            if not record[
                "rationale"
            ]:

                errors.append(
                    (
                        f"{reviewer_name} / {ecr_id}: "
                        "rationale is blank."
                    )
                )


    return errors


# ==========================================================
# JACCARD
# ==========================================================

def jaccard_similarity(
    set_a,
    set_b
):

    union = (
        set_a
        | set_b
    )


    if not union:

        return 1.0


    intersection = (
        set_a
        & set_b
    )


    return (
        len(
            intersection
        )
        /
        len(
            union
        )
    )


# ==========================================================
# FLEISS KAPPA FOR RISK
# ==========================================================

def calculate_fleiss_kappa_risk(
    all_reviews
):
    """
    Fleiss' kappa for categorical risk ratings across
    3 reviewers and 5 ECRs.
    """

    categories = [
        "LOW",
        "MEDIUM",
        "HIGH",
        "CRITICAL"
    ]


    reviewer_names = list(
        all_reviews.keys()
    )


    n_raters = len(
        reviewer_names
    )


    if n_raters < 2:

        return None


    item_category_counts = []


    for ecr_id in EXPECTED_ECRS:

        counter = Counter()


        for reviewer_name in reviewer_names:

            risk = (
                all_reviews[
                    reviewer_name
                ][
                    ecr_id
                ][
                    "risk"
                ]
                .strip()
                .upper()
            )


            counter[
                risk
            ] += 1


        row = [
            counter.get(
                category,
                0
            )
            for category in categories
        ]


        item_category_counts.append(
            row
        )


    # ------------------------------------------------------
    # OBSERVED AGREEMENT
    # ------------------------------------------------------

    item_agreements = []


    for counts in item_category_counts:

        numerator = sum(
            count * (
                count - 1
            )
            for count in counts
        )


        denominator = (
            n_raters
            * (
                n_raters - 1
            )
        )


        item_agreements.append(
            numerator
            / denominator
        )


    p_bar = (
        sum(
            item_agreements
        )
        /
        len(
            item_agreements
        )
    )


    # ------------------------------------------------------
    # EXPECTED AGREEMENT
    # ------------------------------------------------------

    total_ratings = (
        len(
            EXPECTED_ECRS
        )
        * n_raters
    )


    category_totals = [

        sum(
            row[
                category_index
            ]
            for row in item_category_counts
        )

        for category_index in range(
            len(
                categories
            )
        )
    ]


    category_proportions = [

        total
        / total_ratings

        for total in category_totals
    ]


    p_e = sum(
        proportion ** 2
        for proportion in category_proportions
    )


    if p_e == 1:

        kappa = 1.0

    else:

        kappa = (
            p_bar
            - p_e
        ) / (
            1
            - p_e
        )


    return {

        "fleiss_kappa":
            kappa,

        "observed_agreement":
            p_bar,

        "expected_agreement":
            p_e,

        "categories":
            categories
    }


# ==========================================================
# ANALYSE ONE ECR
# ==========================================================

def analyse_ecr(
    ecr_id,
    all_reviews
):

    a = all_reviews[
        "Reviewer_A"
    ][
        ecr_id
    ]


    b = all_reviews[
        "Reviewer_B"
    ][
        ecr_id
    ]


    c = all_reviews[
        "Reviewer_C"
    ][
        ecr_id
    ]


    set_a = a[
        "entities"
    ]


    set_b = b[
        "entities"
    ]


    set_c = c[
        "entities"
    ]


    union = (
        set_a
        | set_b
        | set_c
    )


    unanimous = (
        set_a
        & set_b
        & set_c
    )


    entity_votes = {}


    for entity in sorted(
        union
    ):

        votes = (

            int(
                entity in set_a
            )
            +
            int(
                entity in set_b
            )
            +
            int(
                entity in set_c
            )
        )


        entity_votes[
            entity
        ] = votes


    majority = {

        entity
        for entity, votes
        in entity_votes.items()
        if votes >= 2
    }


    disputed = {

        entity
        for entity, votes
        in entity_votes.items()
        if votes in [
            1,
            2
        ]
    }


    single_reviewer_only = {

        entity
        for entity, votes
        in entity_votes.items()
        if votes == 1
    }


    exact_set_agreement = (
        set_a
        == set_b
        == set_c
    )


    risk_a = a[
        "risk"
    ].strip().upper()


    risk_b = b[
        "risk"
    ].strip().upper()


    risk_c = c[
        "risk"
    ].strip().upper()


    risk_exact_agreement = (
        risk_a
        == risk_b
        == risk_c
    )


    pairwise = {

        "A_B":
            jaccard_similarity(
                set_a,
                set_b
            ),

        "A_C":
            jaccard_similarity(
                set_a,
                set_c
            ),

        "B_C":
            jaccard_similarity(
                set_b,
                set_c
            )
    }


    mean_pairwise = statistics.mean(
        pairwise.values()
    )


    median_pairwise = statistics.median(
        pairwise.values()
    )


    if union:

        unanimous_fraction = (
            len(
                unanimous
            )
            /
            len(
                union
            )
        )

        majority_fraction = (
            len(
                majority
            )
            /
            len(
                union
            )
        )

    else:

        unanimous_fraction = 1.0
        majority_fraction = 1.0


    return {

        "ecr_id":
            ecr_id,

        "reviewer_entities": {

            "Reviewer_A":
                sorted(
                    set_a
                ),

            "Reviewer_B":
                sorted(
                    set_b
                ),

            "Reviewer_C":
                sorted(
                    set_c
                )
        },

        "reviewer_counts": {

            "Reviewer_A":
                len(
                    set_a
                ),

            "Reviewer_B":
                len(
                    set_b
                ),

            "Reviewer_C":
                len(
                    set_c
                )
        },

        "pairwise_jaccard":
            pairwise,

        "mean_pairwise_jaccard":
            mean_pairwise,

        "median_pairwise_jaccard":
            median_pairwise,

        "exact_set_agreement":
            exact_set_agreement,

        "union_entities":
            sorted(
                union
            ),

        "unanimous_entities":
            sorted(
                unanimous
            ),

        "majority_entities":
            sorted(
                majority
            ),

        "disputed_entities":
            sorted(
                disputed
            ),

        "single_reviewer_entities":
            sorted(
                single_reviewer_only
            ),

        "entity_votes":
            entity_votes,

        "unanimous_fraction_of_union":
            unanimous_fraction,

        "majority_fraction_of_union":
            majority_fraction,

        "risk": {

            "Reviewer_A":
                risk_a,

            "Reviewer_B":
                risk_b,

            "Reviewer_C":
                risk_c,

            "exact_agreement":
                risk_exact_agreement
        }
    }


# ==========================================================
# ANALYSE ALL CASES
# ==========================================================

def analyse_all(
    all_reviews
):

    case_results = []


    for ecr_id in EXPECTED_ECRS:

        case_results.append(
            analyse_ecr(
                ecr_id,
                all_reviews
            )
        )


    all_pairwise_scores = []


    for case in case_results:

        all_pairwise_scores.extend(
            case[
                "pairwise_jaccard"
            ].values()
        )


    exact_entity_cases = sum(

        1
        for case in case_results
        if case[
            "exact_set_agreement"
        ]
    )


    exact_risk_cases = sum(

        1
        for case in case_results
        if case[
            "risk"
        ][
            "exact_agreement"
        ]
    )


    total_disputed_entities = sum(

        len(
            case[
                "disputed_entities"
            ]
        )

        for case in case_results
    )


    risk_kappa = (
        calculate_fleiss_kappa_risk(
            all_reviews
        )
    )


    return {

        "case_results":
            case_results,

        "aggregate": {

            "number_of_cases":
                len(
                    case_results
                ),

            "mean_pairwise_jaccard":
                statistics.mean(
                    all_pairwise_scores
                ),

            "median_pairwise_jaccard":
                statistics.median(
                    all_pairwise_scores
                ),

            "minimum_pairwise_jaccard":
                min(
                    all_pairwise_scores
                ),

            "maximum_pairwise_jaccard":
                max(
                    all_pairwise_scores
                ),

            "exact_entity_set_agreement_cases":
                exact_entity_cases,

            "exact_entity_set_agreement_rate":
                (
                    exact_entity_cases
                    / len(
                        case_results
                    )
                ),

            "exact_risk_agreement_cases":
                exact_risk_cases,

            "exact_risk_agreement_rate":
                (
                    exact_risk_cases
                    / len(
                        case_results
                    )
                ),

            "total_disputed_entity_decisions":
                total_disputed_entities,

            "risk_fleiss_kappa":
                risk_kappa
        }
    }


# ==========================================================
# SAVE SUMMARY CSV
# ==========================================================

def save_summary_csv(
    analysis
):

    columns = [

        "ECR_ID",

        "Reviewer_A_Count",
        "Reviewer_B_Count",
        "Reviewer_C_Count",

        "Jaccard_A_B",
        "Jaccard_A_C",
        "Jaccard_B_C",

        "Mean_Jaccard",
        "Median_Jaccard",

        "Exact_Entity_Set_Agreement",

        "Union_Count",
        "Unanimous_Count",
        "Majority_Count",
        "Disputed_Count",

        "Unanimous_Fraction",
        "Majority_Fraction",

        "Risk_A",
        "Risk_B",
        "Risk_C",

        "Exact_Risk_Agreement"
    ]


    with open(
        SUMMARY_CSV,
        "w",
        newline="",
        encoding="utf-8"
    ) as file:

        writer = csv.DictWriter(
            file,
            fieldnames=columns
        )


        writer.writeheader()


        for case in analysis[
            "case_results"
        ]:

            writer.writerow(
                {

                    "ECR_ID":
                        case[
                            "ecr_id"
                        ],

                    "Reviewer_A_Count":
                        case[
                            "reviewer_counts"
                        ][
                            "Reviewer_A"
                        ],

                    "Reviewer_B_Count":
                        case[
                            "reviewer_counts"
                        ][
                            "Reviewer_B"
                        ],

                    "Reviewer_C_Count":
                        case[
                            "reviewer_counts"
                        ][
                            "Reviewer_C"
                        ],

                    "Jaccard_A_B":
                        round(
                            case[
                                "pairwise_jaccard"
                            ][
                                "A_B"
                            ],
                            6
                        ),

                    "Jaccard_A_C":
                        round(
                            case[
                                "pairwise_jaccard"
                            ][
                                "A_C"
                            ],
                            6
                        ),

                    "Jaccard_B_C":
                        round(
                            case[
                                "pairwise_jaccard"
                            ][
                                "B_C"
                            ],
                            6
                        ),

                    "Mean_Jaccard":
                        round(
                            case[
                                "mean_pairwise_jaccard"
                            ],
                            6
                        ),

                    "Median_Jaccard":
                        round(
                            case[
                                "median_pairwise_jaccard"
                            ],
                            6
                        ),

                    "Exact_Entity_Set_Agreement":
                        case[
                            "exact_set_agreement"
                        ],

                    "Union_Count":
                        len(
                            case[
                                "union_entities"
                            ]
                        ),

                    "Unanimous_Count":
                        len(
                            case[
                                "unanimous_entities"
                            ]
                        ),

                    "Majority_Count":
                        len(
                            case[
                                "majority_entities"
                            ]
                        ),

                    "Disputed_Count":
                        len(
                            case[
                                "disputed_entities"
                            ]
                        ),

                    "Unanimous_Fraction":
                        round(
                            case[
                                "unanimous_fraction_of_union"
                            ],
                            6
                        ),

                    "Majority_Fraction":
                        round(
                            case[
                                "majority_fraction_of_union"
                            ],
                            6
                        ),

                    "Risk_A":
                        case[
                            "risk"
                        ][
                            "Reviewer_A"
                        ],

                    "Risk_B":
                        case[
                            "risk"
                        ][
                            "Reviewer_B"
                        ],

                    "Risk_C":
                        case[
                            "risk"
                        ][
                            "Reviewer_C"
                        ],

                    "Exact_Risk_Agreement":
                        case[
                            "risk"
                        ][
                            "exact_agreement"
                        ]
                }
            )


# ==========================================================
# SAVE ENTITY DISAGREEMENT CSV
# ==========================================================

def save_entity_csv(
    analysis
):

    columns = [

        "ECR_ID",
        "Entity_ID",

        "Reviewer_A",
        "Reviewer_B",
        "Reviewer_C",

        "Votes",

        "Agreement_Status"
    ]


    with open(
        ENTITY_CSV,
        "w",
        newline="",
        encoding="utf-8"
    ) as file:

        writer = csv.DictWriter(
            file,
            fieldnames=columns
        )


        writer.writeheader()


        for case in analysis[
            "case_results"
        ]:

            set_a = set(
                case[
                    "reviewer_entities"
                ][
                    "Reviewer_A"
                ]
            )


            set_b = set(
                case[
                    "reviewer_entities"
                ][
                    "Reviewer_B"
                ]
            )


            set_c = set(
                case[
                    "reviewer_entities"
                ][
                    "Reviewer_C"
                ]
            )


            for entity in case[
                "union_entities"
            ]:

                votes = case[
                    "entity_votes"
                ][
                    entity
                ]


                if votes == 3:

                    agreement_status = (
                        "Unanimous Include"
                    )


                elif votes == 2:

                    agreement_status = (
                        "Majority Include - Adjudicate"
                    )


                else:

                    agreement_status = (
                        "Single Reviewer - Adjudicate"
                    )


                writer.writerow(
                    {

                        "ECR_ID":
                            case[
                                "ecr_id"
                            ],

                        "Entity_ID":
                            entity,

                        "Reviewer_A":
                            entity
                            in set_a,

                        "Reviewer_B":
                            entity
                            in set_b,

                        "Reviewer_C":
                            entity
                            in set_c,

                        "Votes":
                            votes,

                        "Agreement_Status":
                            agreement_status
                    }
                )


# ==========================================================
# SAVE JSON
# ==========================================================

def save_json(
    analysis
):

    payload = {

        "phase":
            "5E.1B",

        "analysis_name":
            (
                "Prospective Expert "
                "Inter-Rater Agreement"
            ),

        "created_at":
            datetime.now().isoformat(),

        "reviewers": [
            "Reviewer_A",
            "Reviewer_B",
            "Reviewer_C"
        ],

        "ecrs":
            EXPECTED_ECRS,

        "agent_predictions_accessed":
            False,

        "consensus_ground_truth_created":
            False,

        "analysis":
            analysis
    }


    with open(
        ANALYSIS_JSON,
        "w",
        encoding="utf-8"
    ) as file:

        json.dump(
            payload,
            file,
            indent=2,
            ensure_ascii=False
        )


# ==========================================================
# ADD WORKBOOK SHEET
# ==========================================================

def add_table_sheet(
    workbook,
    title,
    headers,
    rows
):

    worksheet = workbook.create_sheet(
        title
    )


    header_fill = PatternFill(
        fill_type="solid",
        fgColor="D9EAF7"
    )


    thin_border = Border(

        left=Side(
            style="thin",
            color="BFBFBF"
        ),

        right=Side(
            style="thin",
            color="BFBFBF"
        ),

        top=Side(
            style="thin",
            color="BFBFBF"
        ),

        bottom=Side(
            style="thin",
            color="BFBFBF"
        )
    )


    for column, header in enumerate(
        headers,
        start=1
    ):

        cell = worksheet.cell(
            row=1,
            column=column,
            value=header
        )


        cell.font = Font(
            bold=True
        )


        cell.fill = header_fill


        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )


        cell.border = thin_border


    for row_number, row_values in enumerate(
        rows,
        start=2
    ):

        for column, value in enumerate(
            row_values,
            start=1
        ):

            cell = worksheet.cell(
                row=row_number,
                column=column,
                value=value
            )


            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True
            )


            cell.border = thin_border


    worksheet.freeze_panes = "A2"


    # Generic widths.
    for column in range(
        1,
        len(
            headers
        ) + 1
    ):

        worksheet.column_dimensions[
            worksheet.cell(
                row=1,
                column=column
            ).column_letter
        ].width = 22


    return worksheet


# ==========================================================
# CREATE ADJUDICATION WORKBOOK
# ==========================================================

def create_adjudication_workbook(
    analysis,
    all_reviews
):

    if CONSENSUS_TEMPLATE.exists():

        workbook = load_workbook(
            CONSENSUS_TEMPLATE
        )

    else:

        workbook = Workbook()


        if (
            "Sheet"
            in workbook.sheetnames
        ):

            workbook.remove(
                workbook[
                    "Sheet"
                ]
            )


    # Remove old generated analysis sheets
    # if script is rerun.
    for sheet_name in [

        "Agreement_Summary",
        "Entity_Adjudication",
        "Risk_Adjudication",
        "Reviewer_Rationales"

    ]:

        if sheet_name in workbook.sheetnames:

            workbook.remove(
                workbook[
                    sheet_name
                ]
            )


    # ======================================================
    # AGREEMENT SUMMARY
    # ======================================================

    summary_rows = []


    for case in analysis[
        "case_results"
    ]:

        summary_rows.append(
            [

                case[
                    "ecr_id"
                ],

                case[
                    "reviewer_counts"
                ][
                    "Reviewer_A"
                ],

                case[
                    "reviewer_counts"
                ][
                    "Reviewer_B"
                ],

                case[
                    "reviewer_counts"
                ][
                    "Reviewer_C"
                ],

                round(
                    case[
                        "pairwise_jaccard"
                    ][
                        "A_B"
                    ],
                    4
                ),

                round(
                    case[
                        "pairwise_jaccard"
                    ][
                        "A_C"
                    ],
                    4
                ),

                round(
                    case[
                        "pairwise_jaccard"
                    ][
                        "B_C"
                    ],
                    4
                ),

                round(
                    case[
                        "mean_pairwise_jaccard"
                    ],
                    4
                ),

                len(
                    case[
                        "unanimous_entities"
                    ]
                ),

                len(
                    case[
                        "disputed_entities"
                    ]
                ),

                (
                    "Yes"
                    if case[
                        "exact_set_agreement"
                    ]
                    else "No"
                )
            ]
        )


    add_table_sheet(

        workbook,

        "Agreement_Summary",

        [

            "ECR_ID",

            "Reviewer_A_Count",
            "Reviewer_B_Count",
            "Reviewer_C_Count",

            "Jaccard_A_B",
            "Jaccard_A_C",
            "Jaccard_B_C",

            "Mean_Jaccard",

            "Unanimous_Entities",
            "Disputed_Entities",

            "Exact_Set_Agreement"
        ],

        summary_rows
    )


    # ======================================================
    # ENTITY ADJUDICATION
    # ======================================================

    entity_rows = []


    for case in analysis[
        "case_results"
    ]:

        set_a = set(
            case[
                "reviewer_entities"
            ][
                "Reviewer_A"
            ]
        )


        set_b = set(
            case[
                "reviewer_entities"
            ][
                "Reviewer_B"
            ]
        )


        set_c = set(
            case[
                "reviewer_entities"
            ][
                "Reviewer_C"
            ]
        )


        for entity in case[
            "union_entities"
        ]:

            votes = case[
                "entity_votes"
            ][
                entity
            ]


            if votes == 3:

                recommendation = (
                    "Unanimous Include"
                )


            elif votes == 2:

                recommendation = (
                    "Majority Include - Review"
                )


            else:

                recommendation = (
                    "Single Reviewer - Review"
                )


            entity_rows.append(
                [

                    case[
                        "ecr_id"
                    ],

                    entity,

                    (
                        "Include"
                        if entity in set_a
                        else "Exclude"
                    ),

                    (
                        "Include"
                        if entity in set_b
                        else "Exclude"
                    ),

                    (
                        "Include"
                        if entity in set_c
                        else "Exclude"
                    ),

                    votes,

                    recommendation,

                    "",

                    "",

                    "Pending"
                ]
            )


    entity_sheet = add_table_sheet(

        workbook,

        "Entity_Adjudication",

        [

            "ECR_ID",
            "Entity_ID",

            "Reviewer_A",
            "Reviewer_B",
            "Reviewer_C",

            "Votes",

            "Pre_Adjudication_Status",

            "Consensus_Decision",

            "Consensus_Rationale",

            "Adjudication_Status"
        ],

        entity_rows
    )


    entity_sheet.column_dimensions[
        "H"
    ].width = 24


    entity_sheet.column_dimensions[
        "I"
    ].width = 60


    # Highlight human consensus columns.
    consensus_fill = PatternFill(
        fill_type="solid",
        fgColor="FFF2CC"
    )


    for row in range(
        2,
        entity_sheet.max_row + 1
    ):

        for column in [
            8,
            9,
            10
        ]:

            entity_sheet.cell(
                row=row,
                column=column
            ).fill = consensus_fill


    # ======================================================
    # RISK ADJUDICATION
    # ======================================================

    risk_rows = []


    for case in analysis[
        "case_results"
    ]:

        risks = case[
            "risk"
        ]


        risk_counter = Counter(
            [

                risks[
                    "Reviewer_A"
                ],

                risks[
                    "Reviewer_B"
                ],

                risks[
                    "Reviewer_C"
                ]
            ]
        )


        majority_risk = (
            risk_counter
            .most_common(
                1
            )[
                0
            ][
                0
            ]
        )


        majority_votes = (
            risk_counter
            .most_common(
                1
            )[
                0
            ][
                1
            ]
        )


        if majority_votes == 3:

            pre_status = (
                "Unanimous"
            )


        elif majority_votes == 2:

            pre_status = (
                "Majority - Review"
            )


        else:

            pre_status = (
                "No Majority - Review"
            )


        risk_rows.append(
            [

                case[
                    "ecr_id"
                ],

                risks[
                    "Reviewer_A"
                ],

                risks[
                    "Reviewer_B"
                ],

                risks[
                    "Reviewer_C"
                ],

                majority_risk,

                majority_votes,

                pre_status,

                "",

                "",

                "Pending"
            ]
        )


    risk_sheet = add_table_sheet(

        workbook,

        "Risk_Adjudication",

        [

            "ECR_ID",

            "Reviewer_A_Risk",
            "Reviewer_B_Risk",
            "Reviewer_C_Risk",

            "Pre_Adjudication_Majority",

            "Majority_Votes",

            "Pre_Adjudication_Status",

            "Consensus_Risk",

            "Consensus_Rationale",

            "Adjudication_Status"
        ],

        risk_rows
    )


    risk_sheet.column_dimensions[
        "I"
    ].width = 60


    for row in range(
        2,
        risk_sheet.max_row + 1
    ):

        for column in [
            8,
            9,
            10
        ]:

            risk_sheet.cell(
                row=row,
                column=column
            ).fill = consensus_fill


    # ======================================================
    # REVIEWER RATIONALES
    # ======================================================

    rationale_rows = []


    for ecr_id in EXPECTED_ECRS:

        rationale_rows.append(
            [

                ecr_id,

                all_reviews[
                    "Reviewer_A"
                ][
                    ecr_id
                ][
                    "rationale"
                ],

                all_reviews[
                    "Reviewer_B"
                ][
                    ecr_id
                ][
                    "rationale"
                ],

                all_reviews[
                    "Reviewer_C"
                ][
                    ecr_id
                ][
                    "rationale"
                ]
            ]
        )


    rationale_sheet = add_table_sheet(

        workbook,

        "Reviewer_Rationales",

        [

            "ECR_ID",
            "Reviewer_A_Rationale",
            "Reviewer_B_Rationale",
            "Reviewer_C_Rationale"
        ],

        rationale_rows
    )


    rationale_sheet.column_dimensions[
        "B"
    ].width = 65


    rationale_sheet.column_dimensions[
        "C"
    ].width = 65


    rationale_sheet.column_dimensions[
        "D"
    ].width = 65


    workbook.save(
        ADJUDICATION_WORKBOOK
    )


    workbook.close()


# ==========================================================
# PRINT CASE SUMMARY
# ==========================================================

def print_case_summary(
    analysis
):

    print(
        "\n"
        + "=" * 105
    )


    print(
        "PER-ECR INTER-RATER AGREEMENT"
    )


    print(
        "=" * 105
    )


    header = (

        f"{'ECR':<10}"
        f"{'A':>5}"
        f"{'B':>5}"
        f"{'C':>5}"
        f"{'J(A,B)':>10}"
        f"{'J(A,C)':>10}"
        f"{'J(B,C)':>10}"
        f"{'Mean J':>10}"
        f"{'Union':>8}"
        f"{'Unanim':>8}"
        f"{'Dispute':>9}"
    )


    print(
        "\n"
        + header
    )


    print(
        "-" * len(
            header
        )
    )


    for case in analysis[
        "case_results"
    ]:

        print(

            f"{case['ecr_id']:<10}"

            f"{case['reviewer_counts']['Reviewer_A']:>5}"

            f"{case['reviewer_counts']['Reviewer_B']:>5}"

            f"{case['reviewer_counts']['Reviewer_C']:>5}"

            f"{case['pairwise_jaccard']['A_B']:>10.4f}"

            f"{case['pairwise_jaccard']['A_C']:>10.4f}"

            f"{case['pairwise_jaccard']['B_C']:>10.4f}"

            f"{case['mean_pairwise_jaccard']:>10.4f}"

            f"{len(case['union_entities']):>8}"

            f"{len(case['unanimous_entities']):>8}"

            f"{len(case['disputed_entities']):>9}"
        )


# ==========================================================
# MAIN
# ==========================================================

if __name__ == "__main__":

    print(
        "\n"
        + "=" * 100
    )


    print(
        "AUREON AI-PLM AGENT"
    )


    print(
        "Phase 5E.1B - Inter-Rater Agreement Analysis"
    )


    print(
        "=" * 100
    )


    # ======================================================
    # LOAD REVIEWERS
    # ======================================================

    all_reviews = {}


    print(
        "\nLoading reviewer workbooks..."
    )


    for reviewer_name, file_path in (
        REVIEWER_FILES.items()
    ):

        all_reviews[
            reviewer_name
        ] = load_reviewer(
            reviewer_name,
            file_path
        )


        print(
            f"  {reviewer_name}: loaded"
        )


    # ======================================================
    # VALIDATE COMPLETION
    # ======================================================

    completion_errors = (
        validate_reviews_complete(
            all_reviews
        )
    )


    if completion_errors:

        print(
            "\n"
            + "=" * 100
        )


        print(
            "REVIEW COMPLETION CHECK FAILED"
        )


        print(
            "=" * 100
        )


        print(
            "\nThe agreement analysis has NOT been run."
        )


        print(
            "\nResolve the following reviewer items:"
        )


        for error in completion_errors:

            print(
                f"  - {error}"
            )


        print(
            "\nNo consensus or agreement result was created."
        )


        raise SystemExit(
            1
        )


    print(
        "\nAll reviewer completion checks PASSED."
    )


    # ======================================================
    # ANALYSE
    # ======================================================

    analysis = analyse_all(
        all_reviews
    )


    print_case_summary(
        analysis
    )


    aggregate = analysis[
        "aggregate"
    ]


    print(
        "\n"
        + "=" * 100
    )


    print(
        "AGGREGATE AGREEMENT"
    )


    print(
        "=" * 100
    )


    print(
        f"\nMean pairwise Jaccard:   "
        f"{aggregate['mean_pairwise_jaccard']:.4f}"
    )


    print(
        f"Median pairwise Jaccard: "
        f"{aggregate['median_pairwise_jaccard']:.4f}"
    )


    print(
        f"Minimum pairwise Jaccard:"
        f" {aggregate['minimum_pairwise_jaccard']:.4f}"
    )


    print(
        f"Maximum pairwise Jaccard:"
        f" {aggregate['maximum_pairwise_jaccard']:.4f}"
    )


    print(
        "\nExact entity-set agreement:"
    )


    print(
        f"  "
        f"{aggregate['exact_entity_set_agreement_cases']}"
        f"/{aggregate['number_of_cases']} cases"
    )


    print(
        "\nExact risk agreement:"
    )


    print(
        f"  "
        f"{aggregate['exact_risk_agreement_cases']}"
        f"/{aggregate['number_of_cases']} cases"
    )


    print(
        f"\nTotal disputed entity decisions: "
        f"{aggregate['total_disputed_entity_decisions']}"
    )


    risk_kappa = aggregate[
        "risk_fleiss_kappa"
    ]


    if risk_kappa is not None:

        print(
            "\nRisk Fleiss' kappa:"
        )


        print(
            f"  "
            f"{risk_kappa['fleiss_kappa']:.4f}"
        )


        print(
            f"  Observed agreement: "
            f"{risk_kappa['observed_agreement']:.4f}"
        )


        print(
            f"  Expected agreement: "
            f"{risk_kappa['expected_agreement']:.4f}"
        )


    # ======================================================
    # SAVE OUTPUTS
    # ======================================================

    save_summary_csv(
        analysis
    )


    save_entity_csv(
        analysis
    )


    save_json(
        analysis
    )


    create_adjudication_workbook(
        analysis,
        all_reviews
    )


    print(
        "\n"
        + "=" * 100
    )


    print(
        "PHASE 5E.1B OUTPUTS"
    )


    print(
        "=" * 100
    )


    print(
        "\nAgreement summary CSV:"
    )


    print(
        SUMMARY_CSV
    )


    print(
        "\nEntity disagreement CSV:"
    )


    print(
        ENTITY_CSV
    )


    print(
        "\nComplete agreement JSON:"
    )


    print(
        ANALYSIS_JSON
    )


    print(
        "\nConsensus adjudication workbook:"
    )


    print(
        ADJUDICATION_WORKBOOK
    )


    print(
        "\nIMPORTANT:"
    )


    print(
        "These are PRE-CONSENSUS agreement results."
    )


    print(
        "Do NOT run D1 or D2.2 on the five prospective "
        "ECRs yet."
    )


    print(
        "\nNext phase:"
    )


    print(
        "Phase 5E.1C - Expert Consensus Adjudication "
        "and Ground-Truth Freezing"
    )


    print(
        "\n"
        + "=" * 100
    )


    print(
        "PHASE 5E.1B COMPLETE"
    )


    print(
        "=" * 100
    )